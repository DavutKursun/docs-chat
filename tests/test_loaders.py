from __future__ import annotations

from pathlib import Path

import pytest

from docs_chat.loaders import load_file, load_folder, normalize_whitespace, split_into_pseudo_pages


def test_normalize_whitespace_keeps_paragraph_breaks():
    raw = "Line one   with   spaces.\r\n\r\n\r\n\r\nLine two.\t\tTabbed.  \n"
    assert normalize_whitespace(raw) == "Line one with spaces.\n\nLine two. Tabbed."


def test_pdf_pages_keep_real_page_numbers(two_page_pdf: Path):
    pages = load_file(two_page_pdf)
    assert [p.page_no for p in pages] == [1, 2]
    assert "twenty vacation days" in pages[0].text
    assert "twelve months" in pages[1].text
    assert not pages[0].is_pseudo_page


def test_txt_and_md_become_pseudo_pages(tmp_path: Path):
    text = "\n\n".join(f"Paragraph {i} " + "word " * 120 for i in range(10))
    f = tmp_path / "notes.md"
    f.write_text(text, encoding="utf-8")
    pages = load_file(f)
    assert len(pages) >= 3
    assert all(p.is_pseudo_page for p in pages)
    assert [p.page_no for p in pages] == list(range(1, len(pages) + 1))
    assert all(len(p.text.split()) <= 400 + 125 for p in pages)  # ~PSEUDO_PAGE_WORDS, paragraph granularity


def test_docx_is_loaded(tmp_path: Path):
    import docx

    d = docx.Document()
    d.add_paragraph("Remote work is allowed three days per week.")
    d.add_paragraph("Tuesday is the office day.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Budget"
    table.rows[0].cells[1].text = "1200 EUR"
    path = tmp_path / "policy.docx"
    d.save(str(path))
    pages = load_file(path)
    assert len(pages) == 1
    assert "three days per week" in pages[0].text
    assert "Budget | 1200 EUR" in pages[0].text


def test_unsupported_type_raises(tmp_path: Path):
    f = tmp_path / "data.xlsx"
    f.write_bytes(b"not really")
    with pytest.raises(ValueError):
        load_file(f)


def test_scanned_pdf_pages_are_flagged(tmp_path: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.rect(100, 100, 300, 300)  # a drawing, no text
    c.showPage()
    c.save()
    pages, warnings = load_folder(tmp_path)
    assert len(pages) == 1 and pages[0].is_empty
    assert any("no extractable text" in w for w in warnings)


def test_page_furniture_is_stripped(sample_docs_dir: Path):
    from docs_chat.loaders import strip_page_furniture

    assert strip_page_furniture("Page 3\nReal text\n12 of 40\n7 / 9\n") == "Real text"
    assert strip_page_furniture("Chapter 3 starts here") == "Chapter 3 starts here"
    pages = load_file(sample_docs_dir / "aria_x200_manual.pdf")
    assert not pages[1].text.startswith("Page 2")


def test_pseudo_pages_of_empty_text():
    pages = split_into_pseudo_pages("", "empty.txt")
    assert len(pages) == 1 and pages[0].is_empty


def test_sample_docs_load(sample_docs_dir: Path):
    pages, warnings = load_folder(sample_docs_dir)
    assert warnings == []
    names = {p.doc_name for p in pages}
    assert names == {"lumen_employee_handbook.pdf", "aria_x200_manual.pdf"}
    assert all(not p.is_empty for p in pages)
